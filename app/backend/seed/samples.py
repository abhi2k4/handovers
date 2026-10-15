"""Sample upload files for the demo (synthetic), written to graph/seed/northwind/samples/.

- tools-inventory.xlsx   an access inventory (tool, owner, role, renewal): mapped by a fixed template
- zoya-account-notes.docx  free-text account notes for the CS manager: read by MarkItDown + extraction
- resume-backend.txt / resume-marketer.txt  candidate resumes for the screening checker

Run:  uv run python -m seed.samples
"""

from docx import Document
from openpyxl import Workbook

from core.config import get_settings


def main() -> None:
    out = get_settings().graph_dir / "seed" / "northwind" / "samples"
    out.mkdir(parents=True, exist_ok=True)

    wb = Workbook()
    ws = wb.active
    assert ws is not None
    ws.title = "Tools"
    ws.append(["Tool", "Owner", "Role", "Renewal"])
    for row in [
        ["HubSpot", "dev@northwind.example", "admin", "2027-01-31"],
        ["Zendesk", "zoya@northwind.example", "admin", "2026-12-15"],
        ["Figma", "tara@northwind.example", "admin", "2027-03-01"],
        ["Tally", "rahul@northwind.example", "owner", "2027-04-01"],
    ]:
        ws.append(row)
    wb.save(out / "tools-inventory.xlsx")

    doc = Document()
    doc.add_heading("Account notes — Zoya Ali", level=1)
    for line in [
        "Zoya owns the quarterly business reviews for our top 10 merchants and runs them in the first week of each quarter.",
        "She manages the renewal for Bluefin Foods; their contact is Meera Joshi, Head of Finance, who prefers WhatsApp for quick questions.",
        "Zoya handles escalations when a merchant reports a reconciliation mismatch and coordinates the fix with engineering.",
        "Decision (Aug 2026): we stopped offering free onboarding calls to merchants under ₹5L monthly volume, because they rarely converted to paid plans.",
    ]:
        doc.add_paragraph(line)
    doc.save(str(out / "zoya-account-notes.docx"))

    (out / "resume-backend.txt").write_text(
        "Kiran Rao — Backend Engineer\n"
        "4 years building payment APIs in Python (Django, FastAPI) on PostgreSQL.\n"
        "Owned a settlement reconciliation job processing 2M transactions a day; on-call rotation with Sentry and Grafana.\n"
        "Integrated two bank APIs with webhooks and idempotency keys. Wrote Alembic migrations for large tables.\n"
        "Also: Docker, AWS, some Go.\n",
        encoding="utf-8",
    )
    (out / "resume-marketer.txt").write_text(
        "Riya Kapoor — Growth Marketer\n"
        "3 years running Facebook Ads and Instagram Ads for D2C brands; managed ₹8L/month budgets against CAC targets.\n"
        "Built weekly dashboards in Looker Studio from Google Analytics data. Comfortable with Excel.\n"
        "Wrote briefs for creative agencies. Skills: Meta Ads, Canva, SEO basics.\n",
        encoding="utf-8",
    )
    print(f"wrote samples to {out}")


if __name__ == "__main__":
    main()
